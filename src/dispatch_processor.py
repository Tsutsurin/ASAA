"""Обработка Excel-файлов на отработку."""

import logging
import time
from pathlib import Path

from openpyxl import load_workbook

from .ad_user_service import resolve_dispatch_recipients
from .blacklist_service import (
    is_blocked,
    load_blacklist,
)
from .config import BASE_DIR, Settings
from .exchange_sender import send_html_email
from .group_report_splitter import split_report_by_group
from .html_template import render_html_template
from .service_routing import (
    get_service_recipients,
    load_service_routing,
)
from .ticket_number_service import (
    copy_report_to_ticket_folder,
    create_ticket_folder,
    get_max_ticket_number,
)
from .ticket_registry import append_ticket_to_registry
from .utils import join_emails, normalize_text

logger = logging.getLogger(
    'auto_responsible.dispatch'
)


_SENSITIVE_COLUMNS = [
    'Члены группы',
    'Наименование ИС',
    'Ответственный ИС / Администратор ИС',
    'Маршрутизация ИС',
    'Ручной ввод',
    'Бэкэнд сервер',
]


def _build_ticket_subject(
    base: str,
    number: int,
) -> str:
    base = base.rstrip()

    if base.endswith('№'):
        return f'{base}{number}'

    return f'{base} №{number}'


def _wait_before_next(
    settings: Settings,
) -> None:
    pause = (
        settings.dispatch
        .pause_between_emails_seconds
    )

    if pause > 0:
        time.sleep(pause)


def _build_cc(
    resolved_cc: str | None,
    default_cc: str | None,
) -> str | None:
    if resolved_cc is None:
        return default_cc

    return resolved_cc or None


def _remove_attachment_columns(
    file_path: Path,
    names: list[str],
) -> None:
    if not names:
        return

    targets = {
        normalize_text(name)
        for name in names
        if normalize_text(name)
    }

    if not targets:
        return

    wb = load_workbook(file_path)
    ws = wb.active

    columns_to_delete = [
        col_idx
        for col_idx in range(
            1,
            ws.max_column + 1,
        )
        if normalize_text(
            ws.cell(
                row=1,
                column=col_idx,
            ).value
        ) in targets
    ]

    if not columns_to_delete:
        wb.close()
        return

    for col_idx in sorted(
        columns_to_delete,
        reverse=True,
    ):
        ws.delete_cols(col_idx)

    wb.save(file_path)
    wb.close()

    logger.info(
        'Из вложения удалены служебные колонки: %s',
        names,
    )


def _filter_blocked(
    raw_emails: str,
    blacklist: list[str],
    context: str,
) -> list[str]:
    if not raw_emails:
        return []

    normalized = join_emails(
        raw_emails
    )

    if not normalized:
        return []

    allowed = []

    for email in normalized.split('; '):
        email = email.strip()

        if not email:
            continue

        if is_blocked(
            email=email,
            blacklist=blacklist,
        ):
            logger.warning(
                'Получатель заблокирован: %s | %s',
                email,
                context,
            )
            continue

        allowed.append(email)

    return list(
        dict.fromkeys(allowed)
    )


def _filter_and_limit(
    raw_emails: str,
    settings: Settings,
    blacklist: list[str],
    context: str,
) -> list[str]:
    allowed = _filter_blocked(
        raw_emails=raw_emails,
        blacklist=blacklist,
        context=context,
    )

    max_recipients = (
        settings.dispatch
        .recipient_filter
        .max_recipients
    )

    if (
        max_recipients > 0
        and len(allowed) > max_recipients
    ):
        logger.warning(
            'Получателей больше лимита: %s → %s | %s',
            len(allowed),
            max_recipients,
            context,
        )

        allowed = allowed[
            :max_recipients
        ]

    return allowed


def _resolve_recipients(
    group_report,
    settings: Settings,
    service_routing: dict[str, list[str]],
    blacklist: list[str],
) -> tuple[
    str,
    str | None,
    Path,
    str,
    str,
]:
    normal_template = (
        BASE_DIR
        / settings.dispatch.template_path
    )

    service_template = (
        BASE_DIR
        / settings.dispatch.service_template_path
    )

    # 1. Ручной ввод.
    if group_report.manual_emails:
        recipients = _filter_and_limit(
            raw_emails=group_report.manual_emails,
            settings=settings,
            blacklist=blacklist,
            context=(
                'manual | '
                f'{group_report.display_name}'
            ),
        )

        if recipients:
            return (
                '; '.join(recipients),
                None,
                normal_template,
                group_report.display_name,
                'manual',
            )

    # 2. Маршрутизация ИС,
    # записанная непосредственно в Excel.
    if group_report.service_routing_emails:
        recipients = _filter_and_limit(
            raw_emails=(
                group_report
                .service_routing_emails
            ),
            settings=settings,
            blacklist=blacklist,
            context=(
                'service_routing_column | '
                f'{group_report.service_name}'
            ),
        )

        if recipients:
            return (
                '; '.join(recipients),
                None,
                service_template,
                group_report.service_name,
                'service_routing_column',
            )

    # 3. Маршрутизация по service_routing.json.
    if (
        service_routing
        and group_report.service_name
    ):
        json_recipients = (
            get_service_recipients(
                service_name=(
                    group_report.service_name
                ),
                routing=service_routing,
            )
        )

        if json_recipients:
            recipients = _filter_and_limit(
                raw_emails='; '.join(
                    json_recipients
                ),
                settings=settings,
                blacklist=blacklist,
                context=(
                    'service_routing_json | '
                    f'{group_report.service_name}'
                ),
            )

            if recipients:
                return (
                    '; '.join(recipients),
                    None,
                    service_template,
                    group_report.service_name,
                    'service_routing_json',
                )

    # 4. Ответственный ИС /
    # Администратор ИС.
    if group_report.service_owner_admin:
        recipients = _filter_and_limit(
            raw_emails=(
                group_report
                .service_owner_admin
            ),
            settings=settings,
            blacklist=blacklist,
            context=(
                'service_owner_admin | '
                f'{group_report.service_name}'
            ),
        )

        if recipients:
            return (
                '; '.join(recipients),
                None,
                service_template,
                group_report.service_name,
                'service_owner_admin',
            )

        # Если адреса были, но все были исключены
        # blacklist, не используем fallback_to.
        # Переходим к членам группы.
        logger.warning(
            'Все Ответственные ИС / Администраторы ИС '
            'исключены черным списком. '
            'Переход к Членам группы | '
            'service=%s | group=%s',
            group_report.service_name,
            group_report.display_name,
        )

    # 5. Члены группы с фильтрацией через AD.
    if group_report.group_member_emails:
        to, cc = resolve_dispatch_recipients(
            raw_emails=(
                group_report
                .group_member_emails
            ),
            default_to=settings.dispatch.to,
            settings=(
                settings.dispatch
                .recipient_filter
            ),
        )

        return (
            to,
            cc,
            normal_template,
            group_report.display_name,
            'ad_filter',
        )

    # 6. Старое поле Почты / fallback.
    to, cc = resolve_dispatch_recipients(
        raw_emails=group_report.emails,
        default_to=settings.dispatch.to,
        settings=(
            settings.dispatch
            .recipient_filter
        ),
    )

    return (
        to,
        cc,
        normal_template,
        group_report.display_name,
        'ad_filter_fallback',
    )


def process_dispatch_excel_file(
    settings: Settings,
    account,
    source_file: Path,
    original_file_name: str | None = None,
) -> int:
    """
    Обрабатывает Excel-файл для рассылки.

    source_file:
        рабочая копия файла.

    original_file_name:
        исходное имя файла из папки "Отработать".
        Оно записывается в колонку "Система"
        общего реестра заявок.
    """

    # Сохраняем обратную совместимость на случай,
    # если функция вызывается не из
    # dispatch_directory_service.py.
    if not original_file_name:
        original_file_name = source_file.name

    output_dir = (
        BASE_DIR
        / settings.dispatch.output_dir
        / source_file.stem
    )

    group_reports = split_report_by_group(
        source_file=source_file,
        output_dir=output_dir,
        group_column=(
            settings.dispatch.group_column
        ),
        ip_column=(
            settings.dispatch.ip_column
        ),
        fqdn_column=(
            settings.dispatch.fqdn_column
        ),
        vulnerability_column=(
            settings.dispatch
            .vulnerability_column
        ),
    )

    max_emails = (
        settings.dispatch
        .max_emails_per_run
    )

    if (
        max_emails > 0
        and len(group_reports) > max_emails
    ):
        raise RuntimeError(
            'Количество сформированных писем 'f'({len(group_reports)}) превышает '
            f'лимит за запуск ({max_emails}). '
            'Ни одно письмо не отправлено.'
        )

    tasks_dir = Path(
        settings.dispatch.network_tasks_dir
    )

    registry_file = Path(
        settings.dispatch.registry_file
    )

    next_ticket_number = (
        get_max_ticket_number(tasks_dir)
        + 1
    )

    no_group_template = (
        BASE_DIR
        / settings.dispatch
        .no_group_template_path
    )

    service_routing = load_service_routing(
        settings.dispatch.service_routing
    )

    blacklist = load_blacklist(
        settings.dispatch
        .recipient_filter
        .blacklist_file
    )

    created_count = 0
    failed_count = 0

    logger.info(
        'Обработка файла | '
        'working_file=%s | original_file=%s',
        source_file,
        original_file_name,
    )

    for group_report in group_reports:
        try:
            if group_report.has_group:
                ticket_number = next_ticket_number
                next_ticket_number += 1

                (
                    to,
                    cc,
                    template_path,
                    template_value,
                    route_source,
                ) = _resolve_recipients(
                    group_report=group_report,
                    settings=settings,
                    service_routing=service_routing,
                    blacklist=blacklist,
                )

                logger.info(
                    'Маршрут письма | '
                    'group=%s | service=%s | '
                    'route=%s | template_value=%s | '
                    'to=%s',
                    group_report.group_name,
                    group_report.service_name,
                    route_source,
                    template_value,
                    to,
                )

                html_body = render_html_template(
                    template_path=template_path,
                    placeholder=(
                        settings.dispatch.placeholder
                    ),
                    value=template_value,
                )

                ticket_folder = create_ticket_folder(
                    tasks_dir=tasks_dir,
                    ticket_number=ticket_number,
                )

                _remove_attachment_columns(
                    file_path=group_report.file_path,
                    names=(
                        _SENSITIVE_COLUMNS
                        + settings.dispatch
                        .attachment_drop_columns
                    ),
                )

                report_in_ticket = (
                    copy_report_to_ticket_folder(
                        report_file=(
                            group_report.file_path
                        ),
                        ticket_folder=ticket_folder,
                    )
                )

                message_cc = _build_cc(
                    resolved_cc=cc,
                    default_cc=settings.dispatch.cc,
                )

                send_html_email(
                    account=account,
                    to=to,
                    cc=message_cc,
                    subject=_build_ticket_subject(
                        settings.dispatch.subject,
                        ticket_number,
                    ),
                    html_body=html_body,
                    attachments=[
                        group_report.file_path
                    ],
                )

                append_ticket_to_registry(
                    registry_file=registry_file,
                    registry_columns=(
                        settings.dispatch
                        .registry_columns
                    ),
                    ticket_number=ticket_number,
                    group_name=(
                        group_report.display_name
                    ),
                    ip_or_fqdn=(
                        group_report.ip_or_fqdn
                    ),
                    vulnerability_id=(
                        group_report.vulnerability_ids
                    ),
                    status=(
                        settings.dispatch.status_value
                    ),
                    system=original_file_name,
                )

                logger.info(
                    'Заявка %s отправлена | '
                    'system=%s | '
                    'group=%s | service=%s | '
                    'route=%s | to=%s | '
                    'cc=%s | file=%s',
                    ticket_number,
                    original_file_name,
                    group_report.display_name,
                    group_report.service_name,
                    route_source,
                    to,
                    message_cc,
                    report_in_ticket,
                )

            else:
                html_body = render_html_template(
                    template_path=no_group_template,
                    placeholder=(
                        settings.dispatch.placeholder
                    ),
                    value=(
                        group_report.display_name
                    ),
                )

                _remove_attachment_columns(
                    file_path=group_report.file_path,
                    names=(
                        _SENSITIVE_COLUMNS
                        + settings.dispatch
                        .attachment_drop_columns
                    ),
                )

                send_html_email(
                    account=account,
                    to=(
                        settings.dispatch.no_group_to
                    ),
                    cc=(
                        settings.dispatch.no_group_cc
                    ),
                    subject=(
                        settings.dispatch
                        .no_group_subject
                    ),
                    html_body=html_body,
                    attachments=[
                        group_report.file_path
                    ],
                )

                logger.info(
                    'Письмо без группы отправлено'
                )

            created_count += 1

            _wait_before_next(settings)

        except Exception:
            failed_count += 1

            logger.exception(
                'Ошибка обработки группы: %s',
                group_report.display_name,
            )

    if failed_count > 0:
        raise RuntimeError(
            f'Ошибки отправки: {failed_count}. '
            f'Успешно: {created_count}. '
            'Исходный файл не архивируется.'
        )

    logger.info(
        'Файл обработан: %s | '
        'original_file=%s | писем=%s',
        source_file,
        original_file_name,
        created_count,
    )

    return created_count