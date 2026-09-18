"""Работа с общим реестром заявок."""

import logging
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from .utils import normalize_text

logger = logging.getLogger(
    'auto_responsible.ticket_registry'
)


def append_ticket_to_registry(
    registry_file: Path,
    registry_columns: dict[str, str],
    ticket_number: int,
    group_name: str,
    ip_or_fqdn: str,
    vulnerability_id: str,
    status: str,
    system: str,
) -> None:
    """
    Добавляет новую заявку в общий реестр.

    system:
        Оригинальное имя исходного файла
        из директории "Отработать",
        например:
        vulnerabilities_17.09.2026.xlsx
    """

    if not registry_file.exists():
        raise FileNotFoundError(
            'Файл общего реестра '
            f'не найден: {registry_file}'
        )

    wb = load_workbook(
        registry_file
    )

    ws = wb.active

    headers = {}

    for column_index in range(
        1,
        ws.max_column + 1,
    ):
        value = ws.cell(
            row=1,
            column=column_index,
        ).value

        normalized = normalize_text(
            value
        )

        if normalized:
            headers[normalized] = (
                column_index
            )

    values = {
        normalize_text(
            registry_columns[
                'ticket_number'
            ]
        ): ticket_number,

        normalize_text(
            registry_columns[
                'system'
            ]
        ): system,

        normalize_text(
            registry_columns[
                'group'
            ]
        ): group_name,

        normalize_text(
            registry_columns[
                'host'
            ]
        ): ip_or_fqdn,

        normalize_text(
            registry_columns[
                'date'
            ]
        ): datetime.now().strftime(
            '%d.%m.%Y'
        ),

        normalize_text(
            registry_columns[
                'status'
            ]
        ): status,
    }

    vulnerability_column = (
        registry_columns.get(
            'vulnerability_id'
        )
    )

    if vulnerability_column:
        values[
            normalize_text(
                vulnerability_column
            )
        ] = vulnerability_id

    missing_columns = [
        column_name
        for column_name in values
        if column_name not in headers
    ]

    if missing_columns:
        wb.close()

        raise ValueError(
            'В общем реестре отсутствуют '
            'обязательные колонки: '
            + ', '.join(
                missing_columns
            )
        )

    row_number = (
        ws.max_row + 1
    )

    for column_name, value in (
        values.items()
    ):
        column_index = headers[
            column_name
        ]

        ws.cell(
            row=row_number,
            column=column_index,
            value=value,
        )

    try:
        wb.save(
            registry_file
        )

    finally:
        wb.close()

    logger.info(
        'Заявка добавлена в общий реестр | '
        'ticket=%s | system=%s | '
        'group=%s | host=%s | '
        'vulnerability=%s | status=%s',
        ticket_number,
        system,
        group_name,
        ip_or_fqdn,
        vulnerability_id,
        status,
    )