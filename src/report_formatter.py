"""Форматирование Excel-отчетов."""

import configparser
import logging
import re
import sys
import time
import uuid
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)

logger = logging.getLogger('auto_responsible.formatter')


if getattr(sys, 'frozen', False):
    BASE_DIR = Path(sys.executable).resolve().parent
else:
    BASE_DIR = Path(__file__).resolve().parent.parent


_formatter_config_path: Path | None = None


HEADER_FONT = Font(
    bold=True,
    color='000000',
)

HEADER_FILL = PatternFill(
    'solid',
    fgColor='A6A6A6',
)

BORDER_THIN = Border(
    left=Side(style='thin', color='444444'),
    right=Side(style='thin', color='444444'),
    top=Side(style='thin', color='444444'),
    bottom=Side(style='thin', color='444444'),
)

HEADER_ALIGNMENT = Alignment(
    horizontal='center',
    vertical='top',
    wrap_text=True,
)

BODY_ALIGNMENT = Alignment(
    vertical='top',
    wrap_text=False,
)


SEVERITY_COLORS = {
    'critical': 'F2A6A7',
    'high': 'EFD1A4',
    'medium': 'F1EBAC',
    'low': 'CEF2DD',
    'none': 'FFFFFF',
}

SEVERITY_ALIASES = {
    'critical': 'critical',
    'критический': 'critical',
    'критическая': 'critical',

    'high': 'high',
    'высокий': 'high',
    'высокая': 'high',

    'medium': 'medium',
    'средний': 'medium',
    'средняя': 'medium',

    'low': 'low',
    'низкий': 'low',
    'низкая': 'low',

    'none': 'none',
    'нет': 'none',
    'информационный': 'none',
    'информационная': 'none',
}


SEVERITY_FILLS = {}

COL_MAP = []
DROP_COLS_NORM = set()

DEFAULT_COL_WIDTH = 15.0


CVE_RE = re.compile(
    r'CVE-\d{4}-\d{4,7}',
    re.IGNORECASE,
)

BDU_RE = re.compile(
    r'\b(?:\d{4}-\d{1,6}|\d{4,6})\b'
)


VERDICT_MAP = {
    'fail': 'Не выполнено',
    'success': 'Выполнено',
    'nonapplicable': 'Неприменимо',
    'insufficientdata': 'Недостаточно данных',
}


def set_formatter_config(
    path: str | Path,
) -> None:
    global _formatter_config_path

    path = Path(path)

    if not path.is_absolute():
        path = BASE_DIR / path

    _formatter_config_path = path.resolve()

    logger.info(
        'Форматирование Excel: конфиг=%s',
        _formatter_config_path,
    )


def _get_config_path() -> Path:
    if _formatter_config_path is not None:
        return _formatter_config_path

    config_path = BASE_DIR / 'config' / 'config.ini'

    if config_path.is_file():
        return config_path

    return BASE_DIR / 'config.ini'


def norm_name(value) -> str:
    if value is None:
        return ''

    return re.sub(
        r'\s+',
        '',
        str(value),
    ).lower()


def normalize_severity(value) -> str:
    value = str(value or '').strip().lower()

    return SEVERITY_ALIASES.get(
        value,
        value,
    )


def to_blank_space(value):
    if value is None:
        return ' '

    if isinstance(value, str) and value == '':
        return ' '

    return value


def replace_file_with_retry(
    src: Path,
    dst: Path,
    retries: int = 5,
) -> None:
    last_error = None

    for attempt in range(1, retries + 1):
        try:
            src.replace(dst)
            return

        except PermissionError as error:
            last_error = error

            logger.warning(
                'Файл занят, попытка замены %s/%s: %s',
                attempt,
                retries,
                dst,
            )

            time.sleep(2)

    if last_error is not None:
        raise last_error


def load_settings_from_config() -> None:
    global HEADER_FILL
    global SEVERITY_FILLS
    global COL_MAP
    global DROP_COLS_NORM

    config_path = _get_config_path()

    logger.info(
        'Проверяю config.ini: %s | exists=%s',
        config_path,
        config_path.is_file(),
    )

    config = configparser.ConfigParser()
    config.optionxform = str

    if config_path.is_file():
        config.read(
            config_path,
            encoding='utf-8',
        )

        if config.has_section('header'):
            header_color = config.get(
                'header',
                'color',
                fallback=None,
            )

            if header_color:
                HEADER_FILL = PatternFill(
                    'solid',
                    fgColor=header_color.strip(),
                )

        if config.has_section('colors'):
            for key, value in config['colors'].items():
                key = key.strip().lower()
                value = value.strip()

                if key and value:
                    SEVERITY_COLORS[key] = value

        COL_MAP.clear()

        if config.has_section('columns'):
            raw_columns = config.get(
                'columns',
                'cols',
                fallback='',
            )

            for line in raw_columns.splitlines():
                line = line.strip()

                if not line:
                    continue

                if line.endswith(','):
                    line = line[:-1].strip()

                parts = [
                    part.strip()
                    for part in line.split('|')
                ]

                if len(parts) < 2:
                    continue

                source_name = parts[0]
                target_name = parts[1]

                width = None

                if len(parts) >= 3 and parts[2]:
                    try:
                        width = float(parts[2])
                    except ValueError:
                        width = None

                COL_MAP.append(
                    (
                        source_name,
                        target_name,
                        width,
                    )
                )

        DROP_COLS_NORM.clear()

        if config.has_section('drop_columns'):
            raw_drops = config.get(
                'drop_columns',
                'cols',
                fallback='',
            )

            for line in raw_drops.splitlines():
                line = line.strip()

                if not line:
                    continue

                if line.endswith(','):
                    line = line[:-1].strip()

                if line:
                    DROP_COLS_NORM.add(
                        norm_name(line)
                    )

    else:
        logger.warning(
            'config.ini не найден. '
            'Используются встроенные цвета и ширины.'
        )

    SEVERITY_FILLS.clear()

    for key, color in SEVERITY_COLORS.items():
        SEVERITY_FILLS[key.lower()] = PatternFill(
            'solid',
            fgColor=color,
        )


def clean_cve(text: str) -> list[str]:
    if not text:
        return []

    result = []
    seen = set()

    for match in CVE_RE.findall(str(text)):
        value = match.upper()

        if value in seen:
            continue

        seen.add(value)
        result.append(value)

    return result


def clean_bdu(text: str) -> list[str]:
    if not text:
        return []

    result = []
    seen = set()

    for match in BDU_RE.findall(str(text)):
        if match in seen:
            continue

        seen.add(match)
        result.append(match)

    return result


def build_column_maps() -> tuple[
    dict[str, str],
    dict[str, float],
]:
    rename_map = {}
    width_map = {}

    for source_name, target_name, width in COL_MAP:
        source_key = norm_name(source_name)
        target_key = norm_name(target_name)

        rename_map[source_key] = target_name
        rename_map[target_key] = target_name

        if width is not None:
            width_map[source_key] = width
            width_map[target_key] = width

    return rename_map, width_map


def find_column_position(
    source_header: list[str],
    output_columns: list[tuple[int, str]],
    target_name: str,
    source_contains: str | None = None,
) -> int | None:
    target_norm = norm_name(target_name)

    for position, (
        source_index,
        output_name,
    ) in enumerate(output_columns):
        source_name = source_header[source_index]

        source_norm = norm_name(source_name)
        output_norm = norm_name(output_name)

        if output_norm == target_norm:
            return position

        if source_norm == target_norm:
            return position

        if (
            source_contains
            and source_contains.lower() in source_norm
        ):
            return position

    return None


def find_severity_position(
    source_header: list[str],
    output_columns: list[tuple[int, str]],
) -> int | None:
    target = norm_name('Уровень опасности')

    for position, (
        source_index,
        output_name,
    ) in enumerate(output_columns):
        source_name = source_header[source_index]

        source_norm = norm_name(source_name)
        output_norm = norm_name(output_name)

        if output_norm == target:
            return position

        if source_norm == target:
            return position

        if 'severityrating' in source_norm:
            return position

    return None


def stream_transform(
    infile,
    outfile=None,
    progress_callback=None,
) -> Path:
    load_settings_from_config()

    infile = Path(infile)
    outfile = (
        Path(outfile)
        if outfile
        else infile
    )

    logger.info(
        'Начинаю форматирование отчета: %s',
        infile,
    )

    temp_outfile = outfile.with_name(
        f'~formatted_{uuid.uuid4().hex}.xlsx'
    )

    source_wb = load_workbook(
        infile,
        read_only=True,
        data_only=True,
    )

    try:
        source_ws = source_wb.active

        total_rows = max(
            source_ws.max_row - 1,
            0,
        )

        if progress_callback:
            progress_callback(0.0)

        first_row = next(
            source_ws.iter_rows(
                min_row=1,
                max_row=1,
            )
        )
        source_header = [
            (
                str(cell.value).strip()
                if cell.value is not None
                else ''
            )
            for cell in first_row
        ]

        rename_map, width_map = build_column_maps()

        output_columns: list[tuple[int, str]] = []

        for source_index, source_name in enumerate(
            source_header
        ):
            normalized = norm_name(source_name)

            if normalized in DROP_COLS_NORM:
                continue

            # Защита от старых пустых / Unnamed-колонок.
            if not normalized:
                continue

            if normalized.startswith('unnamed:'):
                continue

            output_name = rename_map.get(
                normalized,
                source_name,
            )

            output_columns.append(
                (
                    source_index,
                    output_name,
                )
            )

        wb = Workbook()
        ws = wb.active

        display_header = []

        for _, output_name in output_columns:
            text = str(output_name)

            if ' ' in text:
                first, rest = text.split(
                    ' ',
                    1,
                )

                text = first + '\n' + rest

            display_header.append(text)

        ws.append(display_header)

        for cell in ws[1]:
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = HEADER_ALIGNMENT
            cell.border = BORDER_THIN

        cve_position = find_column_position(
            source_header=source_header,
            output_columns=output_columns,
            target_name='CVE',
            source_contains='cve',
        )

        bdu_position = find_column_position(
            source_header=source_header,
            output_columns=output_columns,
            target_name='BDU',
            source_contains='vulners.ids',
        )

        severity_position = find_severity_position(
            source_header=source_header,
            output_columns=output_columns,
        )

        verdict_position = find_column_position(
            source_header=source_header,
            output_columns=output_columns,
            target_name='Вердикт требования',
        )

        port_position = find_column_position(
            source_header=source_header,
            output_columns=output_columns,
            target_name='Порт',
        )

        logger.info(
            'Колонки форматирования: '
            'severity=%s | cve=%s | '
            'bdu=%s | verdict=%s | port=%s',
            severity_position,
            cve_position,
            bdu_position,
            verdict_position,
            port_position,
        )

        current_row = 1
        processed_rows = 0

        update_every = (
            100
            if total_rows > 1000
            else 10
        )

        for row_cells in source_ws.iter_rows(
            min_row=2
        ):
            current_row += 1
            processed_rows += 1

            row_values = []

            for source_index, _ in output_columns:
                if source_index < len(row_cells):
                    value = row_cells[
                        source_index
                    ].value
                else:
                    value = None

                row_values.append(
                    to_blank_space(value)
                )

            if (
                cve_position is not None
                and cve_position < len(row_values)
            ):
                cve_ids = clean_cve(
                    str(row_values[cve_position])
                )

                row_values[cve_position] = (
                    ','.join(cve_ids)
                    if cve_ids
                    else ' '
                )

            if (
                bdu_position is not None
                and bdu_position < len(row_values)
            ):
                bdu_ids = clean_bdu(
                    str(row_values[bdu_position])
                )

                row_values[bdu_position] = (
                    ', '.join(
                        f'BDU:{bdu_id}'
                        for bdu_id in bdu_ids
                    )
                    if bdu_ids
                    else ' '
                )

            if (
                port_position is not None
                and port_position < len(row_values)
            ):
                raw_port = row_values[
                    port_position
                ]

                cleaned_port = re.sub(
                    r'[^0-9]',
                    '',
                    str(raw_port),
                )

                row_values[port_position] = (
                    cleaned_port
                    if cleaned_port
                    else ' '
                )

            if (
                verdict_position is not None
                and verdict_position < len(row_values)
            ):
                raw_verdict = str(
                    row_values[verdict_position]
                ).strip()

                mapped = VERDICT_MAP.get(
                    raw_verdict.lower()
                )

                if mapped is not None:
                    row_values[
                        verdict_position
                    ] = mapped

            row_fill = None

            if (
                severity_position is not None
                and severity_position < len(row_values)
            ):
                severity = normalize_severity(
                    row_values[severity_position]
                )

                row_fill = SEVERITY_FILLS.get(
                    severity
                )

            for column_index, value in enumerate(
                row_values,
                start=1,
            ):
                cell = ws.cell(
                    row=current_row,
                    column=column_index,
                    value=to_blank_space(value),
                )

                if row_fill is not None:
                    cell.fill = row_fill

                cell.border = BORDER_THIN
                cell.alignment = BODY_ALIGNMENT

                if (
                    cve_position is not None
                    and column_index - 1
                    == cve_position
                ):
                    cell.number_format = '@'

                if (
                    bdu_position is not None
                    and column_index - 1
                    == bdu_position
                ):
                    cell.number_format = '@'

            if (
                progress_callback
                and total_rows > 0
                and (
                    processed_rows % update_every == 0
                    or processed_rows == total_rows
                )
            ):
                progress_callback(
                    processed_rows
                    / total_rows
                    * 100.0
                )

        for index, (
            source_index,
            output_name,
        ) in enumerate(
            output_columns,
            start=1,
        ):
            column_letter = ws.cell(
                row=1,
                column=index,
            ).column_letter

            source_name = source_header[
                source_index
            ]

            source_norm = norm_name(
                source_name
            )

            output_norm = norm_name(
                output_name
            )

            width = (
                width_map.get(source_norm)
                or width_map.get(output_norm)
                or DEFAULT_COL_WIDTH
            )

            ws.column_dimensions[
                column_letter
            ].width = width

        ws.freeze_panes = 'A2'

        if output_columns:
            last_column_letter = ws.cell(
                row=1,
                column=len(output_columns),
            ).column_letter

            ws.auto_filter.ref = (
                f'A1:{last_column_letter}{ws.max_row}'
            )

        wb.save(temp_outfile)
        wb.close()

        if progress_callback:
            progress_callback(100.0)

    finally:
        source_wb.close()

    replace_file_with_retry(
        src=temp_outfile,
        dst=outfile,
    )

    logger.info(
        'Форматирование завершено: %s',
        outfile,
    )

    return outfile