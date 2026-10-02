"""Режим повторной проверки заявок по результатам перескана."""

import logging
import shutil
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

from .config import Settings
from .report_formatter import (
    get_rescan_status_colors,
    stream_transform,
)
from .utils import normalize_text


logger = logging.getLogger(
    'auto_responsible.rescan'
)


RESULT_DIR_NAME = 'Результат'
ARCHIVE_DIR_NAME = 'Архив'
TEMPLATE_FILE_NAME = 'шаблон.txt'

VULNERABILITY_ID_COLUMN = (
    'Идентификатор уязвимости VM'
)

STATUS_COLUMN = 'Статус устранения'

RESULT_COLUMNS = [
    'IP-адрес',
    'Доменное имя',
    'Имя узла',
    'CVE',
    'CVSS Общая',
    'Уровень опасности',
    'Название уязвимости',
    'Уязвимая сущность',
    'Версия уязвимой сущности',
    'Путь установки',
    'Операционная система',
    'Описание уязвимости',
    'Способ устранения уязвимости',
    VULNERABILITY_ID_COLUMN,
]


def _is_excel_file(
    file_path: Path,
) -> bool:
    if not file_path.is_file():
        return False

    if file_path.name.startswith('~$'):
        return False

    return file_path.suffix.lower() in {
        '.xlsx',
        '.xlsm',
    }


def _get_ticket_directories(
    tasks_dir: Path,
) -> list[Path]:
    if not tasks_dir.exists():
        raise FileNotFoundError(
            'Директория заявок не найдена: '
            f'{tasks_dir}'
        )

    directories = [
        path
        for path in tasks_dir.iterdir()
        if path.is_dir()
        and path.name.isdigit()
    ]

    directories.sort(
        key=lambda path: int(path.name)
    )

    return directories


def _get_excel_files(
    ticket_dir: Path,
) -> list[Path]:
    files = [
        path
        for path in ticket_dir.iterdir()
        if _is_excel_file(path)
    ]

    files.sort(
        key=lambda path: (
            path.stat().st_mtime,
            path.name,
        )
    )

    return files


def _find_column(
    dataframe: pd.DataFrame,
    candidates: list[str],
) -> str | None:
    normalized_columns = {
        normalize_text(column): column
        for column in dataframe.columns
    }

    for candidate in candidates:
        normalized_candidate = (
            normalize_text(candidate)
        )

        if (
            normalized_candidate
            in normalized_columns
        ):
            return normalized_columns[
                normalized_candidate
            ]

    return None


def _clean_value(
    value,
) -> str:
    if pd.isna(value):
        return ''

    value = str(value).strip()

    if value.lower() == 'nan':
        return ''

    return value


def _unique_values(
    values: list[str],
) -> list[str]:
    result = []

    for value in values:
        value = _clean_value(value)

        if not value:
            continue

        if value not in result:
            result.append(value)

    return result


def _quote_values(
    values: list[str],
) -> str:
    return ', '.join(
        f'"{value}"'
        for value in values
    )


def _build_template_text(
    ip_addresses: list[str],
    fqdns: list[str],
) -> str:
    ip_line = '; '.join(
        ip_addresses
    )

    fqdn_line = '; '.join(
        fqdns
    )

    plain_ips = ', '.join(
        ip_addresses
    )

    quoted_fqdns = _quote_values(
        fqdns
    )

    lines = [
        ip_line,
        fqdn_line,
        (
            'Host.IpAddress in '
            f'[{plain_ips}]'
        ),
        (
            'host.fqdn in '
            f'[{quoted_fqdns}]'
        ),
        (
            'WebSite.DomainName in '
            f'[{quoted_fqdns}]'
        ),
    ]

    return '\n'.join(
        lines
    )


def _create_template_if_missing(
    ticket_dir: Path,
    original_file: Path,
) -> None:
    template_file = (
        ticket_dir
        / TEMPLATE_FILE_NAME
    )

    if template_file.exists():
        return

    try:
        dataframe = pd.read_excel(
            original_file
        )

    except Exception:
        logger.exception(
            'Не удалось прочитать Excel '
            'для создания шаблона: %s',
            original_file,
        )
        return

    ip_column = _find_column(
        dataframe,
        [
            'IP-адрес',
            'IP адрес',
            'IP',
            'Host.IpAddress',
            'IP Address',
            'IpAddress',
        ],
    )

    fqdn_column = _find_column(
        dataframe,
        [
            'Доменное имя',
            'FQDN',
            'Host.Fqdn',
            'Hostname',
        ],
    )

    ip_addresses = []

    if ip_column:
        ip_addresses = _unique_values(
            dataframe[
                ip_column
            ].tolist()
        )

    fqdns = []

    if fqdn_column:
        fqdns = _unique_values(
            dataframe[
                fqdn_column
            ].tolist()
        )

    template_text = (
        _build_template_text(
            ip_addresses=ip_addresses,
            fqdns=fqdns,
        )
    )

    template_file.write_text(
        template_text,
        encoding='utf-8',
    )

    logger.info(
        'Создан шаблон: %s | '
        'IP=%s | FQDN=%s',
        template_file,
        len(ip_addresses),
        len(fqdns),
    )


def _ensure_result_directories(
    ticket_dir: Path,
) -> tuple[Path, Path]:
    result_dir = (
        ticket_dir
        / RESULT_DIR_NAME
    )

    archive_dir = (
        result_dir
        / ARCHIVE_DIR_NAME
    )

    result_dir.mkdir(
        parents=True,
        exist_ok=True,
    )

    archive_dir.mkdir(
        parents=True,
        exist_ok=True,
    )

    return (
        result_dir,
        archive_dir,
    )


def _normalize_vulnerability_id(
    value,
) -> str:
    value = _clean_value(
        value
    )

    return value.casefold()


def _get_vulnerability_column(
    dataframe: pd.DataFrame,
) -> str:
    column = _find_column(
        dataframe,
        [
            VULNERABILITY_ID_COLUMN,
        ],
    )

    if not column:
        raise ValueError(
            'В Excel отсутствует колонка '
            f'"{VULNERABILITY_ID_COLUMN}"'
        )

    return column


def _get_original_and_rescan(
    excel_files: list[Path],
) -> tuple[Path, Path]:
    if len(excel_files) != 2:
        raise ValueError(
            'Для сравнения должно быть '
            'ровно 2 Excel-файла. '
            f'Найдено: {len(excel_files)}'
        )

    sorted_files = sorted(
        excel_files,
        key=lambda path: (
            path.stat().st_mtime,
            path.name,
        ),
    )

    return (
        sorted_files[0],
        sorted_files[1],
    )


def _get_result_source_column(
    dataframe: pd.DataFrame,
    target_column: str,
) -> str | None:
    aliases = {
        'IP-адрес': [
            'IP-адрес',
            'IP адрес',
            'IP',
            'Host.IpAddress',
            'IP Address',
            'IpAddress',
        ],
        'Доменное имя': [
            'Доменное имя',
            'FQDN',
            'Host.Fqdn',
            'Hostname',
        ],
        'Имя узла': [
            'Имя узла',
            'HOST.hostname',
            'Hostname',
        ],
        'CVE': [
            'CVE',
            'Q.cve',
        ],
        'CVSS Общая': [
            'CVSS Общая',
            'CVSS',
            'CVSS Score',
        ],
        'Уровень опасности': [
            'Уровень опасности',
            'Severity',
            'Критичность',
        ],
        'Название уязвимости': [
            'Название уязвимости',
            'Name',
        ],
        'Уязвимая сущность': [
            'Уязвимая сущность',
            'Уязвимая сущность (ОС/ПО/Сервис)',
            'VulnerableEntity',
        ],
        'Версия уязвимой сущности': [
            'Версия уязвимой сущности',
            'Version',
        ],
        'Путь установки': [
            'Путь установки',
            'Path',
        ],
        'Операционная система': [
            'Операционная система',
            'OS',
        ],
        'Описание уязвимости': [
            'Описание уязвимости',
            'Описание',
            'Description',
        ],
        'Способ устранения уязвимости': [
            'Способ устранения уязвимости',
            'Способ устранения',
            'Рекомендации',
            'Рекомендация',
            'Remediation',
        ],
        VULNERABILITY_ID_COLUMN: [
            VULNERABILITY_ID_COLUMN,
        ],
    }

    return _find_column(
        dataframe,
        aliases.get(
            target_column,
            [target_column],
        ),
    )


def _build_result_dataframe(
    original_df: pd.DataFrame,
    rescan_df: pd.DataFrame,
) -> pd.DataFrame:
    original_vulnerability_column = (
        _get_vulnerability_column(
            original_df
        )
    )

    rescan_vulnerability_column = (
        _get_vulnerability_column(
            rescan_df
        )
    )

    rescan_ids = {
        _normalize_vulnerability_id(
            value
        )
        for value in rescan_df[
            rescan_vulnerability_column
        ].tolist()
        if _normalize_vulnerability_id(
            value
        )
    }

    source_columns = {
        target_column: (
            _get_result_source_column(
                original_df,
                target_column,
            )
        )
        for target_column
        in RESULT_COLUMNS
    }

    result_rows = []

    for _, row in original_df.iterrows():
        vulnerability_id = (
            _clean_value(
                row.get(
                    original_vulnerability_column,
                    '',
                )
            )
        )

        normalized_id = (
            _normalize_vulnerability_id(
                vulnerability_id
            )
        )

        if (
            normalized_id
            and normalized_id
            in rescan_ids
        ):
            status = 'неустранено'
        else:
            status = 'устранено'

        result_row = {}

        for target_column in RESULT_COLUMNS:
            source_column = (
                source_columns[
                    target_column
                ]
            )

            if source_column:
                result_row[
                    target_column
                ] = _clean_value(
                    row.get(
                        source_column,
                        '',
                    )
                )
            else:
                result_row[
                    target_column
                ] = ''

        result_row[
            STATUS_COLUMN
        ] = status

        result_rows.append(
            result_row
        )

    return pd.DataFrame(
        result_rows,
        columns=(
            RESULT_COLUMNS
            + [STATUS_COLUMN]
        ),
    )


def _format_rescan_status(
    result_file: Path,
) -> None:
    colors = (
        get_rescan_status_colors()
    )

    resolved_fill = PatternFill(
        fill_type='solid',
        fgColor=colors['resolved'],
    )

    unresolved_fill = PatternFill(
        fill_type='solid',
        fgColor=colors['unresolved'],
    )

    wb = load_workbook(
        result_file
    )

    ws = wb.active
    ws.title = 'Проверка'

    status_column_index = None

    for column_index in range(
        1,
        ws.max_column + 1,
    ):
        header = normalize_text(
            ws.cell(
                row=1,
                column=column_index,
            ).value
        )

        if header == normalize_text(
            STATUS_COLUMN
        ):
            status_column_index = (
                column_index
            )
            break

    if status_column_index is None:
        wb.close()

        raise ValueError(
            'Не найдена колонка '
            f'"{STATUS_COLUMN}"'
        )

    rows = list(
        ws.iter_rows(
            min_row=2,
        )
    )

    rows.sort(
        key=lambda row: (
            str(
                row[
                    status_column_index - 1
                ].value
                or ''
            ).casefold()
        )
    )

    values = [
        [
            cell.value
            for cell in row
        ]
        for row in rows
    ]

    for row_index, row_values in enumerate(
        values,
        start=2,
    ):
        status = str(
            row_values[
                status_column_index - 1
            ]
            or ''
        ).strip().casefold()

        if status == 'устранено':
            fill = resolved_fill

        elif status == 'неустранено':
            fill = unresolved_fill

        else:
            fill = None

        for column_index, value in enumerate(
            row_values,
            start=1,
        ):
            cell = ws.cell(
                row=row_index,
                column=column_index,
            )

            cell.value = value

            if fill is not None:
                cell.fill = fill

    wb.save(
        result_file
    )

    wb.close()


def _create_result_file(
    ticket_dir: Path,
    result_dir: Path,
    original_file: Path,
    rescan_file: Path,
) -> Path:
    # Второй Excel является сырым результатом
    # повторного сканирования.
    # Сначала приводим его к общему формату ASAA.
    stream_transform(
        infile=rescan_file,
        outfile=rescan_file,
    )

    original_df = pd.read_excel(
        original_file
    )

    rescan_df = pd.read_excel(
        rescan_file
    )

    result_df = (
        _build_result_dataframe(
            original_df=original_df,
            rescan_df=rescan_df,
        )
    )

    date_string = (
        datetime.now().strftime(
            '%d.%m.%Y'
        )
    )

    result_file = (
        result_dir
        / (
            f'{ticket_dir.name}'
            f'_проверка_'
            f'{date_string}.xlsx'
        )
    )

    if result_file.exists():
        timestamp = (
            datetime.now().strftime(
                '%H%M%S'
            )
        )

        result_file = (
            result_dir
            / (
                f'{ticket_dir.name}'
                f'_проверка_'
                f'{date_string}_'
                f'{timestamp}.xlsx'
            )
        )

    result_df.to_excel(
        result_file,
        index=False,
    )

    # Применяем к результату тот же
    # formatter/config.ini, что используется
    # для остальных Excel ASAA.
    stream_transform(
        infile=result_file,
        outfile=result_file,
    )

    # После общего форматирования заменяем
    # severity-заливку на статусную.
    _format_rescan_status(
        result_file
    )

    unresolved_count = int(
        (
            result_df[
                STATUS_COLUMN
            ]
            == 'неустранено'
        ).sum()
    )

    resolved_count = int(
        (
            result_df[
                STATUS_COLUMN
            ]
            == 'устранено'
        ).sum()
    )

    logger.info(
        'Создан результат перескана | '
        'ticket=%s | original=%s | '
        'rescan=%s | resolved=%s | '
        'unresolved=%s | result=%s',
        ticket_dir.name,
        original_file.name,
        rescan_file.name,
        resolved_count,
        unresolved_count,
        result_file,
    )

    return result_file


def _move_rescan_to_archive(
    rescan_file: Path,
    archive_dir: Path,
) -> Path:
    destination = (
        archive_dir
        / rescan_file.name
    )

    if destination.exists():
        timestamp = (
            datetime.now().strftime(
                '%Y%m%d_%H%M%S'
            )
        )

        destination = (
            archive_dir
            / (
                f'{rescan_file.stem}_'
                f'{timestamp}'
                f'{rescan_file.suffix}'
            )
        )

    shutil.move(
        str(rescan_file),
        str(destination),
    )

    logger.info(
        'Перескан перемещён в архив: %s',
        destination,
    )

    return destination


def _prepare_ticket_directory(
    ticket_dir: Path,
) -> None:
    result_dir, archive_dir = (
        _ensure_result_directories(
            ticket_dir
        )
    )

    excel_files = _get_excel_files(
        ticket_dir
    )

    if not excel_files:
        logger.warning(
            'В заявке %s нет Excel-файлов',
            ticket_dir.name,
        )
        return

    original_file = min(
        excel_files,
        key=lambda path: (
            path.stat().st_mtime,
            path.name,
        ),
    )

    _create_template_if_missing(
        ticket_dir=ticket_dir,
        original_file=original_file,
    )

    if len(excel_files) == 1:
        logger.info(
            'Заявка %s подготовлена. '
            'Перескан пока отсутствует.',
            ticket_dir.name,
        )
        return

    if len(excel_files) > 2:
        logger.error(
            'В заявке %s найдено больше '
            'двух Excel-файлов: %s. '
            'Сравнение пропущено.',
            ticket_dir.name,
            [
                path.name
                for path in excel_files
            ],
        )
        return

    original_file, rescan_file = (
        _get_original_and_rescan(
            excel_files
        )
    )

    _create_result_file(
        ticket_dir=ticket_dir,
        result_dir=result_dir,
        original_file=original_file,
        rescan_file=rescan_file,
    )

    _move_rescan_to_archive(
        rescan_file=rescan_file,
        archive_dir=archive_dir,
    )


def run_rescan_mode(
    settings: Settings,
) -> int:
    tasks_dir = Path(
        settings.dispatch.network_tasks_dir
    )

    ticket_directories = (
        _get_ticket_directories(
            tasks_dir
        )
    )

    logger.info(
        'Запуск режима пересканирования. '
        'Папок заявок: %s',
        len(ticket_directories),
    )

    processed = 0
    failed = 0

    for ticket_dir in ticket_directories:
        try:
            _prepare_ticket_directory(
                ticket_dir
            )

            processed += 1

        except Exception:
            failed += 1

            logger.exception(
                'Ошибка режима перескана '
                'для заявки %s',
                ticket_dir.name,
            )

    logger.info(
        'Режим пересканирования завершён. '
        'Обработано=%s | ошибок=%s',
        processed,
        failed,
    )

    return processed