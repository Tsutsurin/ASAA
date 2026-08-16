"""Модуль чтения Excel, извлечения параметров и построения обогащённых отчётов."""

import copy
import logging
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .utils import normalize_text, normalize_value

logger = logging.getLogger('auto_responsible.excel_parser')


DEFAULT_FQDN_CANDIDATES = [
    'Доменное имя',
    'FQDN',
    'DNS Name',
    'Host.Fqdn',
]

DEFAULT_IP_CANDIDATES = [
    'IP-адрес',
    'IP',
    'IP Address',
    'Host.Ip',
]


def _find_column(df: pd.DataFrame, candidates: list[str]) -> str | None:
    normalized_map = {
        normalize_text(column): column
        for column in df.columns
    }
    for candidate in candidates:
        if normalize_text(candidate) in normalized_map:
            return normalized_map[normalize_text(candidate)]
    return None


def _find_header_column(ws, candidates: list[str]) -> int | None:
    header_map = {
        normalize_text(ws.cell(row=1, column=col_idx).value): col_idx
        for col_idx in range(1, ws.max_column + 1)
    }
    for candidate in candidates:
        key = normalize_text(candidate)
        if key in header_map:
            return header_map[key]
    return None


def extract_params_from_files(
    files: list[Path],
    columns_config: dict[str, list[str]],
) -> pd.DataFrame:
    rows = []

    fqdn_candidates = columns_config.get('fqdn', DEFAULT_FQDN_CANDIDATES)
    ip_candidates = columns_config.get('ip', DEFAULT_IP_CANDIDATES)

    for file_path in files:
        file_path = Path(file_path)
        logger.info('Читаю Excel: %s', file_path)

        df = pd.read_excel(file_path)

        fqdn_column = _find_column(df, fqdn_candidates)
        ip_column = _find_column(df, ip_candidates)

        if not fqdn_column and not ip_column:
            logger.warning('Колонки fqdn/ip не найдены: %s', file_path)
            continue

        for _, row in df.iterrows():
            fqdn = normalize_value(row.get(fqdn_column) if fqdn_column else '')
            ip = normalize_value(row.get(ip_column) if ip_column else '')

            if fqdn or ip:
                rows.append({'source_file': str(file_path), 'fqdn': fqdn, 'ip': ip})

    result = pd.DataFrame(rows)
    logger.info('Извлечено параметров: %s', len(result))
    return result


def _coalesce(row, *keys: str) -> str:
    for key in keys:
        value = row.get(key)
        if value is None or pd.isna(value):
            continue
        text = str(value).strip()
        if text and text.lower() != 'nan':
            return text
    return ''


def normalize_api_row(row) -> dict:
    return {
        **row.to_dict(),
        'group': normalize_value(_coalesce(row, 'group', 'Группа', 'responsible_group', 'responsible')),
        'group_members': normalize_value(_coalesce(row, 'group_members', 'Члены группы', 'emails', 'Почты', 'email', 'mail')),
        'service': normalize_value(_coalesce(row, 'service', 'Наименование ИС', 'Service')),
        'service_owner_admin': normalize_value(_coalesce(row, 'service_owner_admin', 'Ответственный ИС / Администратор ИС')),
        'service_routing': normalize_value(_coalesce(row, 'service_routing', 'Маршрутизация ИС')),
        'backend_fqdn': normalize_value(_coalesce(row, 'backend_fqdn', 'Бэкэнд сервер')),
    }


def build_lookup_map(enriched_params: pd.DataFrame) -> dict[tuple[str, str], dict]:
    result = {}

    for _, row in enriched_params.iterrows():
        normalized = normalize_api_row(row)

        fqdn = normalize_value(_coalesce(row, 'fqdn', 'FQDN', 'Доменное имя')).lower()
        ip = normalize_value(_coalesce(row, 'ip', 'IP', 'IP-адрес'))

        if fqdn or ip:
            result[(fqdn, ip)] = normalized
        if fqdn:
            result[(fqdn, '')] = normalized
        if ip:
            result[('', ip)] = normalized

    logger.info('Lookup сформирован: %s ключей', len(result))
    return result


def _get_last_business_column(ws) -> int:
    col_idx = ws.max_column
    while col_idx > 1:
        empty = all(
            normalize_value(ws.cell(row=r, column=col_idx).value) == ''
            for r in range(1, ws.max_row + 1)
        )
        if not empty:
            break
        col_idx -= 1
    return col_idx


def _get_insert_column(ws) -> int:
    last = _get_last_business_column(ws)
    return last + 1 if last < ws.max_column else ws.max_column + 1


def _copy_column_style(ws, source_col: int, target_col: int) -> None:
    ws.column_dimensions[get_column_letter(target_col)].width = (
        ws.column_dimensions[get_column_letter(source_col)].width
    )
    for row_idx in range(1, ws.max_row + 1):
        src = ws.cell(row=row_idx, column=source_col)
        dst = ws.cell(row=row_idx, column=target_col)
        if src.has_style:
            dst._style = copy.copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        if src.alignment:
            dst.alignment = copy.copy(src.alignment)
        if src.border:
            dst.border = copy.copy(src.border)
        if src.fill:
            dst.fill = copy.copy(src.fill)
        if src.font:
            dst.font = copy.copy(src.font)


def ensure_column(ws, header_name: str) -> int:
    for col_idx in range(1, ws.max_column + 1):
        if normalize_text(ws.cell(row=1, column=col_idx).value) == normalize_text(header_name):
            return col_idx

    insert_col = _get_insert_column(ws)
    ws.insert_cols(insert_col)
    _copy_column_style(ws, max(1, insert_col - 1), insert_col)
    ws.cell(row=1, column=insert_col, value=header_name)
    return insert_col


def _reset_filters(ws) -> None:
    for row_idx in range(1, ws.max_row + 1):
        ws.row_dimensions[row_idx].hidden = False
    last = _get_last_business_column(ws)
    ws.auto_filter.ref = f'A1:{get_column_letter(last)}{ws.max_row}'


def build_output_report(
    source_file: Path,
    enriched_params: pd.DataFrame,
    output_file: Path,
    columns_config: dict[str, list[str]],
) -> None:
    source_file = Path(source_file)
    output_file = Path(output_file)

    logger.info('Формирую отчет: %s', output_file)

    wb = load_workbook(source_file)
    ws = wb.active
    _reset_filters(ws)

    fqdn_candidates = columns_config.get('fqdn', DEFAULT_FQDN_CANDIDATES)
    ip_candidates = columns_config.get('ip', DEFAULT_IP_CANDIDATES)

    fqdn_col_idx = _find_header_column(ws, fqdn_candidates)
    ip_col_idx = _find_header_column(ws, ip_candidates)

    col_indices = {
        'Группа': ensure_column(ws, 'Группа'),
        'Члены группы': ensure_column(ws, 'Члены группы'),
        'Наименование ИС': ensure_column(ws, 'Наименование ИС'),
        'Ответственный ИС / Администратор ИС': ensure_column(
            ws, 'Ответственный ИС / Администратор ИС',
        ),
        'Маршрутизация ИС': ensure_column(ws, 'Маршрутизация ИС'),
        'Ручной ввод': ensure_column(ws, 'Ручной ввод'),
        'Бэкэнд сервер': ensure_column(ws, 'Бэкэнд сервер'),
    }

    lookup_map = build_lookup_map(enriched_params)
    filled_count = 0

    for row_idx in range(2, ws.max_row + 1):
        fqdn = normalize_value(ws.cell(row=row_idx, column=fqdn_col_idx).value) if fqdn_col_idx else ''
        ip = normalize_value(ws.cell(row=row_idx, column=ip_col_idx).value) if ip_col_idx else ''

        enriched_row = (
            lookup_map.get((fqdn.lower(), ip))
            or lookup_map.get((fqdn.lower(), ''))
            or lookup_map.get(('', ip))
        )

        if not enriched_row:
            continue

        ws.cell(row=row_idx, column=col_indices['Группа'], value=enriched_row.get('group'))
        ws.cell(row=row_idx, column=col_indices['Члены группы'], value=enriched_row.get('group_members'))
        ws.cell(row=row_idx, column=col_indices['Наименование ИС'], value=enriched_row.get('service'))
        ws.cell(row=row_idx, column=col_indices['Ответственный ИС / Администратор ИС'], value=enriched_row.get('service_owner_admin'))
        ws.cell(row=row_idx, column=col_indices['Маршрутизация ИС'], value=enriched_row.get('service_routing'))
        ws.cell(row=row_idx, column=col_indices['Бэкэнд сервер'], value=enriched_row.get('backend_fqdn'))

        filled_count += 1

    _reset_filters(ws)
    wb.save(output_file)

    logger.info('Отчет сохранен: %s | заполнено: %s', output_file, filled_count)