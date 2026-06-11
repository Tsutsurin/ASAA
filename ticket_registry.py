import logging
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

logger = logging.getLogger('auto_responsible.ticket_registry')


def normalize_header(value) -> str:
    return (
        str(value or '')
        .strip()
        .replace('\n', ' ')
    )


def get_header_map(ws) -> dict[str, int]:
    result = {}

    for idx, cell in enumerate(ws[1], start=1):
        header = normalize_header(cell.value)

        if header:
            result[header] = idx

    return result


def append_ticket_to_registry(
    registry_file: Path,
    registry_columns: dict[str, str],
    ticket_number: int,
    group_name: str,
    ip_or_fqdn: str,
    vulnerability_id: str,
    status: str,
) -> None:
    if not registry_file.exists():
        raise FileNotFoundError(
            f'Файл реестра не найден: {registry_file}'
        )

    wb = load_workbook(registry_file)
    ws = wb.active

    header_map = get_header_map(ws)

    logger.info(
        'Колонки реестра: %s',
        list(header_map.keys()),
    )

    values = {
        registry_columns['ticket_number']: ticket_number,
        registry_columns['group']: group_name,
        registry_columns['host']: ip_or_fqdn,
        registry_columns['date']: datetime.now().strftime('%d.%m.%Y'),
        registry_columns['status']: status,
    }

    vulnerability_column = registry_columns.get('vulnerability_id')

    if vulnerability_column:
        values[vulnerability_column] = vulnerability_id

    missing_columns = [
        column_name
        for column_name in values.keys()
        if column_name not in header_map
    ]

    if missing_columns:
        raise ValueError(
            'В реестре отсутствуют колонки: '
            f'{missing_columns}'
        )

    row_idx = ws.max_row + 1

    for column_name, value in values.items():
        ws.cell(
            row=row_idx,
            column=header_map[column_name],
            value=value,
        )

    wb.save(registry_file)

    logger.info(
        'Запись добавлена в реестр: '
        'ticket=%s group=%s host=%s vuln_id=%s',
        ticket_number,
        group_name,
        ip_or_fqdn,
        vulnerability_id,
    )