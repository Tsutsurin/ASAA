import logging
import time
from pathlib import Path

from openpyxl import load_workbook

from ad_user_service import (
    email_blocked,
    resolve_dispatch_recipients,
    split_emails,
)
from config import BASE_DIR, Settings
from exchange_client import (
    download_excel_attachments,
    find_latest_email_in_folder,
    mark_as_read,
)
from exchange_sender import send_html_email
from group_report_splitter import split_report_by_group
from html_template import render_html_template
from service_routing import (
    get_service_recipients,
    load_service_routing,
)
from ticket_number_service import (
    copy_report_to_ticket_folder,
    create_ticket_folder,
    get_max_ticket_number,
)
from ticket_registry import append_ticket_to_registry

logger = logging.getLogger('auto_responsible.dispatch')


def build_ticket_subject(base_subject: str, ticket_number: int) -> str:
    if base_subject.endswith('№'):
        return f'{base_subject}{ticket_number}'

    if base_subject.endswith('№ '):
        return f'{base_subject}{ticket_number}'

    return f'{base_subject} №{ticket_number}'


def wait_before_next_email(settings: Settings) -> None:
    pause = settings.dispatch.pause_between_emails_seconds

    if pause > 0:
        time.sleep(pause)


def join_raw_emails(*values: str) -> str:
    result = []

    for value in values:
        result.extend(split_emails(value))

    return '; '.join(dict.fromkeys(result))


def normalize_header(value) -> str:
    return (
        str(value or '')
        .replace('\n', ' ')
        .replace('\r', ' ')
        .replace('\t', ' ')
        .strip()
        .lower()
    )


def remove_attachment_columns(
    file_path: Path,
    column_names: list[str],
) -> None:
    if not column_names:
        return

    targets = {
        normalize_header(name)
        for name in column_names
        if normalize_header(name)
    }

    if not targets:
        return

    wb = load_workbook(file_path)
    ws = wb.active

    columns_to_delete = []

    for col_idx in range(1, ws.max_column + 1):
        header = normalize_header(ws.cell(row=1, column=col_idx).value)

        if header in targets:
            columns_to_delete.append(col_idx)

    if not columns_to_delete:
        wb.close()
        logger.info(
            'Колонки для удаления во вложении не найдены: %s | файл=%s',
            column_names,
            file_path,
        )
        return

    for col_idx in sorted(columns_to_delete, reverse=True):
        ws.delete_cols(col_idx)

    wb.save(file_path)
    wb.close()

    logger.info(
        'Из вложения удалены колонки: %s | файл=%s',
        column_names,
        file_path,
    )


def filter_blocked_and_limit_emails(
    raw_emails: str,
    settings: Settings,
    context: str,
) -> list[str]:
    source_emails = split_emails(raw_emails)

    allowed_emails = []

    for email in source_emails:
        if email_blocked(
            email=email,
            blocked_emails=settings.dispatch.recipient_filter.blocked_emails,
        ):
            logger.warning(
                'Получатель заблокирован черным списком: %s | context=%s',
                email,
                context,
            )
            continue

        allowed_emails.append(email)

    allowed_emails = list(dict.fromkeys(allowed_emails))

    max_recipients = settings.dispatch.recipient_filter.max_recipients

    if max_recipients > 0 and len(allowed_emails) > max_recipients:
        logger.warning(
            'Получателей больше лимита: %s. Оставлено: %s | context=%s',
            len(allowed_emails),
            max_recipients,
            context,
        )
        allowed_emails = allowed_emails[:max_recipients]

    return allowed_emails


def build_message_cc(
    filtered_cc: str | None,
    default_cc: str | None,
) -> str | None:
    if filtered_cc is None:
        return default_cc

    return filtered_cc


def process_dispatch_excel_file(
    settings: Settings,
    account,
    source_file: Path,
) -> int:
    output_dir = (
        BASE_DIR
        / settings.dispatch.output_dir
        / source_file.stem
    )

    group_reports = split_report_by_group(
        source_file=source_file,
        output_dir=output_dir,
        group_column=settings.dispatch.group_column,
        ip_column=settings.dispatch.ip_column,
        fqdn_column=settings.dispatch.fqdn_column,
        vulnerability_column=settings.dispatch.vulnerability_column,
    )

    tasks_dir = Path(settings.dispatch.network_tasks_dir)
    registry_file = Path(settings.dispatch.registry_file)

    next_ticket_number = get_max_ticket_number(tasks_dir) + 1

    normal_template_path = BASE_DIR / settings.dispatch.template_path
    no_group_template_path = BASE_DIR / settings.dispatch.no_group_template_path
    service_template_path = BASE_DIR / settings.dispatch.service_template_path

    created_count = 0
    failed_count = 0
    max_emails = settings.dispatch.max_emails_per_run

    service_routing = load_service_routing(
        settings.dispatch.service_routing,
    )

    for group_report in group_reports:
        if max_emails > 0 and created_count >= max_emails:
            logger.warning(
                'Достигнут лимит писем за запуск: %s. Остальные группы не обработаны.',
                max_emails,
            )
            break

        try:
            if group_report.has_group:
                ticket_number = next_ticket_number
                next_ticket_number += 1

                ticket_folder = create_ticket_folder(
                    tasks_dir=tasks_dir,
                    ticket_number=ticket_number,
                )

                append_ticket_to_registry(
                    registry_file=registry_file,
                    registry_columns=settings.dispatch.registry_columns,
                    ticket_number=ticket_number,
                    group_name=group_report.display_name,
                    ip_or_fqdn=group_report.ip_or_fqdn,
                    vulnerability_id=group_report.vulnerability_ids,
                    status=settings.dispatch.status_value,
                )

                raw_service_recipients = get_service_recipients(
                    service_name=group_report.service_name,
                    routing=service_routing,
                )

                filtered_to = ''
                filtered_cc = None
                template_path = normal_template_path
                template_value = group_report.display_name
                route_source = ''

                service_routing_recipients = filter_blocked_and_limit_emails(
                    raw_emails='; '.join(raw_service_recipients),
                    settings=settings,
                    context=f'service_routing | {group_report.service_name}',
                )

                if service_routing_recipients:
                    filtered_to = '; '.join(service_routing_recipients)
                    filtered_cc = None
                    template_path = service_template_path
                    template_value = group_report.service_name
                    route_source = 'service_routing'

                    logger.info(
                        'Использована маршрутизация по ИС для группы %s | ИС=%s | to=%s | cc=%s',
                        group_report.display_name,
                        group_report.service_name,
                        filtered_to,
                        settings.dispatch.cc,
                    )

                else:
                    serviceusers_raw_emails = join_raw_emails(
                        group_report.service_owner_emails,
                        group_report.service_admin_emails,
                    )

                    serviceusers_recipients = filter_blocked_and_limit_emails(
                        raw_emails=serviceusers_raw_emails,
                        settings=settings,
                        context=f'serviceusers | {group_report.service_name}',
                    )

                    if serviceusers_recipients:
                        filtered_to = '; '.join(serviceusers_recipients)
                        filtered_cc = None
                        template_path = service_template_path
                        template_value = group_report.service_name
                        route_source = 'serviceusers'

                        logger.info(
                            'Использованы Ответственный ИС/Администратор ИС для группы %s | ИС=%s | to=%s | cc=%s',
                            group_report.display_name,
                            group_report.service_name,
                            filtered_to,
                            settings.dispatch.cc,
                        )

                    else:
                        filtered_to, filtered_cc = resolve_dispatch_recipients(
                            raw_emails=group_report.emails,
                            default_to=settings.dispatch.to,
                            settings=settings.dispatch.recipient_filter,
                        )
                        template_path = normal_template_path
                        template_value = group_report.display_name
                        route_source = 'ad_filter'

                        logger.info(
                            'Получатели после AD-фильтрации для группы %s: to=%s | cc=%s',
                            group_report.display_name,
                            filtered_to,
                            build_message_cc(
                                filtered_cc=filtered_cc,
                                default_cc=settings.dispatch.cc,
                            ),
                        )

                html_body = render_html_template(
                    template_path=template_path,
                    placeholder=settings.dispatch.placeholder,
                    group_name=template_value,
                )

                remove_attachment_columns(
                    file_path=group_report.file_path,column_names=settings.dispatch.attachment_drop_columns,
                )

                report_in_ticket_folder = copy_report_to_ticket_folder(
                    report_file=group_report.file_path,
                    ticket_folder=ticket_folder,
                )

                send_html_email(
                    account=account,
                    to=filtered_to,
                    cc=build_message_cc(
                        filtered_cc=filtered_cc,
                        default_cc=settings.dispatch.cc,
                    ),
                    subject=build_ticket_subject(
                        settings.dispatch.subject,
                        ticket_number,
                    ),
                    html_body=html_body,
                    attachments=[group_report.file_path],
                )

                logger.info(
                    'Создана и отправлена заявка %s для группы %s | route=%s | файл в заявке: %s',
                    ticket_number,
                    group_report.display_name,
                    route_source,
                    report_in_ticket_folder,
                )

            else:
                html_body = render_html_template(
                    template_path=no_group_template_path,
                    placeholder=settings.dispatch.placeholder,
                    group_name=group_report.display_name,
                )

                remove_attachment_columns(
                    file_path=group_report.file_path,
                    column_names=settings.dispatch.attachment_drop_columns,
                )

                send_html_email(
                    account=account,
                    to=settings.dispatch.no_group_to,
                    cc=settings.dispatch.no_group_cc,
                    subject=settings.dispatch.no_group_subject,
                    html_body=html_body,
                    attachments=[group_report.file_path],
                )

                logger.info(
                    'Отправлено письмо для строк без группы без номера заявки'
                )

            created_count += 1
            wait_before_next_email(settings)

        except Exception:
            failed_count += 1

            logger.exception(
                'Ошибка отправки/создания заявки для группы: %s',
                group_report.display_name,
            )

    if failed_count > 0:
        raise RuntimeError(
            f'Ошибки при отправке писем: {failed_count}. '
            f'Успешно отправлено: {created_count}. '
            f'Исходный файл не должен архивироваться.'
        )

    logger.info(
        'Обработка Excel-файла на отработку завершена: %s | отправлено писем: %s',
        source_file,
        created_count,
    )

    return created_count


def process_dispatch_folder(settings: Settings, account) -> bool:
    if not settings.dispatch.enabled:
        logger.info('Сценарий Отработать отключен')
        return False

    message = find_latest_email_in_folder(
        account=account,
        folder_name=settings.dispatch_folder,
        only_unread=settings.only_unread,
        subject_contains='[отработать]',
    )

    if message is None:
        logger.info('В папке Отработать нет подходящих писем')
        return False

    logger.info(
        'Найдено письмо на отработку: subject=%s',
        str(message.subject or ''),
    )

    excel_files = download_excel_attachments(message)

    if not excel_files:
        logger.warning('В письме Отработать нет Excel-вложений')
        return False

    source_file = excel_files[0]

    process_dispatch_excel_file(
        settings=settings,
        account=account,
        source_file=source_file,
    )

    mark_as_read(message)

    logger.info('Сценарий Отработать из почты завершен')

    return True