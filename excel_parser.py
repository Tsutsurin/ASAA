import copy
import logging
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

logger = logging.getLogger('auto_responsible.excel_parser')


def normalize_column_name(value) -> str:
    return (
        str(value or '')
        .replace('\n', ' ')
        .replace('\r', ' ')
        .strip()
        .lower()
    )


def normalize_value(value) -> str:
    if value is None:
        return ''

    if pd.isna(value):
        return ''

    value = str(value).strip()

    if value.endswith('.0'):
        value = value[:-2]

    return value


def is_blank_value(value) -> bool:
    return str(value or '').strip() == ''


def is_blank_column(ws, col_idx: int) -> bool:
    for row_idx in range(1, ws.max_row + 1):
        if not is_blank_value(ws.cell(row=row_idx, column=col_idx).value):
            return False

    return True


def get_last_business_column(ws) -> int:
    col_idx = ws.max_column

    while col_idx > 1 and is_blank_column(ws, col_idx):
        col_idx -= 1

    return col_idx


def get_insert_column(ws) -> int:
    last_business_col = get_last_business_column(ws)

    if last_business_col < ws.max_column:
        return last_business_col + 1

    return ws.max_column + 1


def copy_column_style(ws, source_col: int, target_col: int) -> None:
    ws.column_dimensions[get_column_letter(target_col)].width = (
        ws.column_dimensions[get_column_letter(source_col)].width
    )

    for row_idx in range(1, ws.max_row + 1):
        source_cell = ws.cell(row=row_idx, column=source_col)
        target_cell = ws.cell(row=row_idx, column=target_col)

        if source_cell.has_style:
            target_cell._style = copy.copy(source_cell._style)

        if source_cell.number_format:
            target_cell.number_format = source_cell.number_format

        if source_cell.alignment:
            target_cell.alignment = copy.copy(source_cell.alignment)

        if source_cell.border:
            target_cell.border = copy.copy(source_cell.border)

        if source_cell.fill:
            target_cell.fill = copy.copy(source_cell.fill)

        if source_cell.font:
            target_cell.font = copy.copy(source_cell.font)


def find_column(
    df: pd.DataFrame,
    candidates: list[str],
) -> str | None:
    normalized_map = {
        normalize_column_name(column): column
        for column in df.columns
    }

    for candidate in candidates:
        normalized_candidate = normalize_column_name(candidate)

        if normalized_candidate in normalized_map:
            return normalized_map[normalized_candidate]

    return None


def extract_params_from_files(
    files: list[Path],
    columns_config: dict[str, list[str]],
) -> pd.DataFrame:
    rows = []

    fqdn_candidates = columns_config.get(
        'fqdn',
        [
            'Доменное имя',
            'FQDN',
            'DNS Name',
            'Host.Fqdn',
        ],
    )

    ip_candidates = columns_config.get(
        'ip',
        [
            'IP-адрес',
            'IP',
            'IP Address',
            'Host.Ip',
        ],
    )

    for file_path in files:
        file_path = Path(file_path)

        logger.info(
            'Читаю Excel для извлечения параметров: %s',
            file_path,
        )

        df = pd.read_excel(file_path)

        fqdn_column = find_column(
            df=df,
            candidates=fqdn_candidates,
        )

        ip_column = find_column(
            df=df,
            candidates=ip_candidates,
        )

        if not fqdn_column and not ip_column:
            logger.warning(
                'В файле не найдены колонки fqdn/ip: %s',
                file_path,
            )
            continue

        for _, row in df.iterrows():
            fqdn = normalize_value(
                row.get(fqdn_column)
                if fqdn_column
                else ''
            )

            ip = normalize_value(
                row.get(ip_column)
                if ip_column
                else ''
            )

            if not fqdn and not ip:
                continue

            rows.append(
                {
                    'source_file': str(file_path),
                    'fqdn': fqdn,
                    'ip': ip,
                }
            )

    result = pd.DataFrame(rows)

    logger.info(
        'Извлечено параметров для обогащения: %s',
        len(result),
    )

    return result


def normalize_api_row(row) -> dict:
    group = (
        row.get('group')
        or row.get('Группа')
        or row.get('responsible_group')
        or row.get('responsible')
        or ''
    )

    leader_email = (
        row.get('leader_email')
        or row.get('Лидер группы')
        or row.get('leader')
        or row.get('main_email')
        or ''
    )

    emails = (
        row.get('emails')
        or row.get('Почты')
        or row.get('email')
        or row.get('mail')
        or ''
    )

    return {
        **row.to_dict(),
        'group': normalize_value(group),
        'leader_email': normalize_value(leader_email),
        'emails': normalize_value(emails),
    }


def build_lookup_map(
    enriched_params: pd.DataFrame,
) -> dict[tuple[str, str], dict]:
    result = {}

    for _, row in enriched_params.iterrows():
        normalized_row = normalize_api_row(row)

        fqdn = normalize_value(
            row.get('fqdn')
            or row.get('FQDN')
            or row.get('Доменное имя')
        )

        ip = normalize_value(
            row.get('ip')
            or row.get('IP')
            or row.get('IP-адрес')
        )

        fqdn_key = fqdn.lower()
        ip_key = ip

        if fqdn_key or ip_key:
            result[(fqdn_key, ip_key)] = normalized_row

        if fqdn_key:
            result[(fqdn_key, '')] = normalized_row

        if ip_key:
            result[('', ip_key)] = normalized_row

    logger.info(
        'Сформирован lookup для обогащения: %s ключей',
        len(result),
    )

    return result


def find_header_column(ws, candidates: list[str]) -> int | None:
    header_map = {}

    for col_idx in range(1, ws.max_column + 1):
        value = ws.cell(row=1, column=col_idx).value
        normalized = normalize_column_name(value)

        if normalized:
            header_map[normalized] = col_idx

    for candidate in candidates:
        normalized_candidate = normalize_column_name(candidate)

        if normalized_candidate in header_map:
            return header_map[normalized_candidate]

    return None


def ensure_column(ws, header_name: str) -> int:
    for col_idx in range(1, ws.max_column + 1):
        value = ws.cell(row=1, column=col_idx).value

        if normalize_column_name(value) == normalize_column_name(header_name):
            return col_idx

    insert_col = get_insert_column(ws)

    ws.insert_cols(insert_col)

    source_col = max(1, insert_col - 1)
    copy_column_style(
        ws=ws,
        source_col=source_col,
        target_col=insert_col,
    )

    ws.cell(row=1, column=insert_col, value=header_name)

    return insert_col


def reset_filters_and_show_rows(ws) -> None:
    for row_idx in range(1, ws.max_row + 1):
        ws.row_dimensions[row_idx].hidden = False

    last_business_col = get_last_business_column(ws)
    last_letter = get_column_letter(last_business_col)

    ws.auto_filter.ref = f'A1:{last_letter}{ws.max_row}'


def build_output_report(
    source_file: Path,
    enriched_params: pd.DataFrame,
    output_file: Path,
    columns_config: dict[str, list[str]],
) -> None:
    source_file = Path(source_file)
    output_file = Path(output_file)

    logger.info(
        'Начинаю формирование итогового отчета: %s',
        output_file,
    )

    wb = load_workbook(source_file)
    ws = wb.active

    reset_filters_and_show_rows(ws)

    fqdn_candidates = columns_config.get(
        'fqdn',
        [
            'Доменное имя',
            'FQDN',
            'DNS Name',
            'Host.Fqdn',
        ],
    )

    ip_candidates = columns_config.get(
        'ip',
        [
            'IP-адрес',
            'IP',
            'IP Address',
            'Host.Ip',
        ],
    )

    fqdn_col_idx = find_header_column(
        ws=ws,
        candidates=fqdn_candidates,
    )

    ip_col_idx = find_header_column(
        ws=ws,
        candidates=ip_candidates,
    )

    group_col_idx = ensure_column(ws, 'Группа')
    leader_col_idx = ensure_column(ws, 'Лидер группы')
    email_col_idx = ensure_column(ws, 'Почты')

    lookup_map = build_lookup_map(enriched_params)

    filled_count = 0

    for row_idx in range(2, ws.max_row + 1):
        fqdn = ''

        if fqdn_col_idx:
            fqdn = normalize_value(
                ws.cell(
                    row=row_idx,
                    column=fqdn_col_idx,
                ).value
            )

        ip = ''

        if ip_col_idx:
            ip = normalize_value(
                ws.cell(
                    row=row_idx,
                    column=ip_col_idx,
                ).value
            )

        enriched_row = (
            lookup_map.get((fqdn.lower(), ip))
            or lookup_map.get((fqdn.lower(), ''))
            or lookup_map.get(('', ip))
        )

        if not enriched_row:
            continue

        group = normalize_value(enriched_row.get('group'))
        leader = normalize_value(enriched_row.get('leader_email'))
        emails = normalize_value(enriched_row.get('emails'))

        ws.cell(
            row=row_idx,
            column=group_col_idx,
            value=group,
        )

        ws.cell(
            row=row_idx,
            column=leader_col_idx,
            value=leader,
        )

        ws.cell(
            row=row_idx,
            column=email_col_idx,
            value=emails,
        )

        filled_count += 1

    reset_filters_and_show_rows(ws)

    wb.save(output_file)

    logger.info(
        'Итоговый отчет сохрансохранен: %s | заполнено строк: %s',
        output_file,
        filled_count,
    )