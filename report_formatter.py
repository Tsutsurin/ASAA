import configparser
import logging
import re
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

logger = logging.getLogger('auto_responsible.formatter')

BASE_DIR = Path(__file__).resolve().parent
CONFIG_PATH = BASE_DIR / 'config.ini'

HEADER_FONT = Font(bold=True, color='000000')
HEADER_FILL = PatternFill('solid', fgColor='1F2933')

BORDER_THIN = Border(
    left=Side(style='thin', color='444444'),
    right=Side(style='thin', color='444444'),
    top=Side(style='thin', color='444444'),
    bottom=Side(style='thin', color='444444'),
)

BODY_ALIGNMENT = Alignment(vertical='top', wrap_text=False)

SEVERITY_COLORS = {
    'critical': 'f2a6a7',
    'high': 'efd1a4',
    'medium': 'f1ebac',
    'low': 'cef2dd',
    'none': 'FFFFFF',
}

SEVERITY_FILLS = {}

VERDICT_MAP = {
    'fail': 'Не выполнено',
    'success': 'Выполнено',
    'nonapplicable': 'Неприменимо',
    'insufficientdata': 'Недостаточно данных',
}

COL_MAP = []
DROP_COLS_NORM = set()

DEFAULT_COL_WIDTH = 15.0

CVE_RE = re.compile(r'CVE-\d{4}-\d{4,7}', re.IGNORECASE)
BDU_RE = re.compile(r'\b(?:\d{4}-\d{1,6}|\d{4,6})\b')


def norm_name(s: str) -> str:
    if s is None:
        return ''
    return re.sub(r'\s+', '', str(s)).lower()


def to_blank_space(v):
    if v is None:
        return ' '
    if isinstance(v, str) and v == '':
        return ' '
    return v


def load_settings_from_config() -> None:
    global HEADER_FILL, SEVERITY_FILLS
    global COL_MAP, DROP_COLS_NORM

    config = configparser.ConfigParser()
    config.optionxform = str

    if CONFIG_PATH.is_file():
        config.read(CONFIG_PATH, encoding='utf-8')

        if config.has_section('header'):
            header_color = config.get('header', 'color', fallback=None)
            if header_color:
                HEADER_FILL = PatternFill('solid', fgColor=header_color.strip())

        if config.has_section('colors'):
            for key, val in config['colors'].items():
                SEVERITY_COLORS[key.strip().lower()] = val.strip()

        COL_MAP.clear()

        if config.has_section('columns'):
            raw_cols = config.get('columns', 'cols', fallback='')
            for line in raw_cols.splitlines():
                line = line.strip()

                if not line:
                    continue

                if line.endswith(','):
                    line = line[:-1].strip()

                parts = [p.strip() for p in line.split('|')]

                if len(parts) < 2:
                    continue

                src = parts[0]
                dst = parts[1]

                width = None
                if len(parts) >= 3 and parts[2]:
                    try:
                        width = float(parts[2])
                    except ValueError:
                        width = None

                COL_MAP.append((src, dst, width))

        DROP_COLS_NORM.clear()

        if config.has_section('drop_columns'):
            raw_drops = config.get('drop_columns', 'cols', fallback='')
            for line in raw_drops.splitlines():
                line = line.strip()

                if not line:
                    continue

                if line.endswith(','):
                    line = line[:-1].strip()

                DROP_COLS_NORM.add(norm_name(line))

    SEVERITY_FILLS.clear()

    for key, color in SEVERITY_COLORS.items():
        SEVERITY_FILLS[key.lower()] = PatternFill('solid', fgColor=color)


def clean_cve(text: str) -> list[str]:
    if not text:
        return []

    ids = []
    seen = set()

    for match in CVE_RE.findall(str(text)):
        value = match.upper()

        if value not in seen:
            seen.add(value)
            ids.append(value)

    return ids


def clean_bdu(text: str) -> list[str]:
    if not text:
        return []

    ids = []
    seen = set()

    for match in BDU_RE.findall(str(text)):
        if match not in seen:
            seen.add(match)
            ids.append(match)

    return ids


def stream_transform(infile, outfile=None, progress_callback=None) -> Path:
    load_settings_from_config()

    infile = Path(infile)
    outfile = Path(outfile) if outfile else infile

    logger.info('Начинаю форматирование отчета: %s', infile)

    temp_outfile = outfile.with_suffix('.formatted_tmp.xlsx')

    src_wb = load_workbook(infile, read_only=True, data_only=True)

    try:
        src_ws = src_wb.active

        total_rows = max(src_ws.max_row - 1, 0)
        processed_rows = 0

        if progress_callback:
            progress_callback(0.0)

        first_row = next(src_ws.iter_rows(min_row=1, max_row=1))
        src_header = [
            str(cell.value).strip() if cell.value is not None else ''
            for cell in first_row
        ]

        rename_map_norm = {}
        width_map_norm = {}

        for src, dst, width in COL_MAP:
            key = norm_name(src)
            rename_map_norm[key] = dst

            if width is not None:
                width_map_norm[key] = width

        exist_idx_map = []
        tgt_header = []

        for idx, src_name in enumerate(src_header):
            normalized = norm_name(src_name)

            if normalized in DROP_COLS_NORM:
                continue

            out_name = rename_map_norm.get(normalized, src_name)

            exist_idx_map.append((idx, out_name))
            tgt_header.append(out_name)

        exist_idx_map.append((None, ' '))
        tgt_header.append(' ')

        wb = Workbook()
        ws = wb.active

        display_header = []

        for src_idx, out_name in exist_idx_map:
            if src_idx is None:
                display_header.append('')
            else:
                text = str(out_name)

                if ' ' in text:
                    first, rest = text.split(' ', 1)
                    text = first + '\n' + rest

                display_header.append(text)

        ws.append(display_header)

        for col_idx, cell in enumerate(ws[1], start=1):
            src_idx, _out_name = exist_idx_map[col_idx - 1]
            is_blank_col = src_idx is None

            if not is_blank_col:
                cell.font = HEADER_FONT
                cell.fill = HEADER_FILL
                cell.alignment = Alignment(
                    horizontal='center',
                    vertical='top',
                    wrap_text=True,
                )
                cell.border = BORDER_THIN
            else:
                cell.border = Border()

        cve_pos = None
        bdu_pos = None
        sev_pos = None
        verdict_pos = None
        port_pos = None

        sev_target_name_norm = norm_name('Уровень опасности')
        verdict_target_norm = norm_name('Вердикт требования')
        port_target_norm = norm_name('Порт')

        for pos, (src_idx, out_name) in enumerate(exist_idx_map):
            if src_idx is None:
                continue

            src_name = src_header[src_idx]
            src_norm = norm_name(src_name)
            out_norm = norm_name(out_name)

            if cve_pos is None and (out_norm == 'cve' or 'cve' in src_norm):
                cve_pos = pos

            if bdu_pos is None and (
                out_norm == 'bdu'
                or 'vulners.ids' in src_norm
                or 'bdu' in src_norm
            ):
                bdu_pos = pos

            if sev_pos is None and (
                out_norm == sev_target_name_norm
                or 'severityrating' in src_norm
            ):
                sev_pos = pos

            if verdict_pos is None and (
                out_norm == verdict_target_norm
                or src_norm == verdict_target_norm
            ):
                verdict_pos = pos

            if port_pos is None and (
                out_norm == port_target_norm
                or src_norm == port_target_norm
            ):
                port_pos = pos

        current_row = 1
        update_every = 100 if total_rows > 1000 else 10

        for row_cells in src_ws.iter_rows(min_row=2):
            current_row += 1
            processed_rows += 1

            row_vals = []

            for src_idx, _out_name in exist_idx_map:
                if src_idx is None:
                    value = ' '
                else:
                    value = row_cells[src_idx].value if src_idx < len(row_cells) else None
                    value = to_blank_space(value)

                row_vals.append(value)

            if cve_pos is not None and cve_pos < len(row_vals):
                cve_ids = clean_cve(str(row_vals[cve_pos]))
                row_vals[cve_pos] = ', '.join(cve_ids) if cve_ids else ' '

            if bdu_pos is not None and bdu_pos < len(row_vals):
                bdu_ids = clean_bdu(str(row_vals[bdu_pos]))

                if bdu_ids:
                    row_vals[bdu_pos] = ', '.join(
                        f'BDU:{bdu_id}'
                        for bdu_id in bdu_ids
                    )
                else:
                    row_vals[bdu_pos] = ' '

            if port_pos is not None and port_pos < len(row_vals):
                raw_port = row_vals[port_pos]
                cleaned_port = re.sub(r'[^0-9]', '', str(raw_port))
                row_vals[port_pos] = cleaned_port if cleaned_port else ' '

            if verdict_pos is not None and verdict_pos < len(row_vals):
                raw = str(row_vals[verdict_pos]).strip()
                key = raw.lower()
                mapped = VERDICT_MAP.get(key)

                if mapped is not None:
                    row_vals[verdict_pos] = mapped

            row_fill = None

            if sev_pos is not None and sev_pos < len(row_vals):
                sev_raw = str(row_vals[sev_pos]).strip().lower()
                row_fill = SEVERITY_FILLS.get(sev_raw)

            for col_idx, ((src_idx, _out_name),value) in enumerate(
                zip(exist_idx_map, row_vals),
                start=1,
            ):
                value = to_blank_space(value)

                cell = ws.cell(
                    row=current_row,
                    column=col_idx,
                    value=value,
                )

                is_blank_col = src_idx is None

                if row_fill is not None and not is_blank_col:
                    cell.fill = row_fill

                if not is_blank_col:
                    cell.border = BORDER_THIN
                else:
                    cell.border = Border()

                cell.alignment = BODY_ALIGNMENT

                if not is_blank_col and (
                    (cve_pos is not None and col_idx - 1 == cve_pos)
                    or (bdu_pos is not None and col_idx - 1 == bdu_pos)
                ):
                    cell.number_format = '@'

            if progress_callback and total_rows > 0 and (
                processed_rows % update_every == 0
                or processed_rows == total_rows
            ):
                progress = processed_rows / total_rows * 100.0
                progress_callback(progress)

        for index, (src_idx, _out_name) in enumerate(exist_idx_map, start=1):
            col_letter = ws.cell(row=1, column=index).column_letter

            if src_idx is None:
                width = DEFAULT_COL_WIDTH
            else:
                src_name = src_header[src_idx]
                normalized = norm_name(src_name)
                width = width_map_norm.get(normalized, DEFAULT_COL_WIDTH)

            ws.column_dimensions[col_letter].width = width

        ws.freeze_panes = 'A2'

        last_row = ws.max_row
        last_filter_col_idx = max(1, len(tgt_header) - 1)
        last_col_letter = ws.cell(row=1, column=last_filter_col_idx).column_letter
        ws.auto_filter.ref = f'A1:{last_col_letter}{last_row}'

        wb.save(temp_outfile)

        if progress_callback:
            progress_callback(100.0)

    finally:
        src_wb.close()

    temp_outfile.replace(outfile)

    logger.info('Форматирование завершено: %s', outfile)

    return outfile